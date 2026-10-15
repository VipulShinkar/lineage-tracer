import hashlib
import io

import psycopg
import pytest
from fastapi.testclient import TestClient

from lineage_store import audit, db, impact, snapshot
from lineage_store.store import LineageStore

CELL = "C07.00:Corporates:r0010:c0200"
ELEMENT = "field:collateral_master.collateral_value"
SUP = {"X-Role": "supervisor", "X-Actor": "Supervisor [A]"}
ALL = ["cover", "narrative", "diagram", "sourcing", "transformations", "waterfall", "adjustments", "controls",
       "contributors", "limitations"]


@pytest.fixture()
def st(conn):
    return LineageStore(conn)


@pytest.fixture()
def client(lineage_dsn, tmp_path, monkeypatch):
    from lineage_api.app import app

    monkeypatch.setenv("LINEAGE_EVIDENCE_DIR", str(tmp_path))
    return TestClient(app)


def test_coverage_seed_fix(st):
    cov = {c: v for c, v in st.q("SELECT cell_id, coverage FROM cell_def WHERE coverage <> 'not_in_scope'")}
    assert cov["C07.00:Corporates:r0010:c0220"] == "complete"
    assert cov["C07.00:Institutions:r0010:c0200"] == "partial" and cov[CELL] == "partial"
    assert list(cov.values()).count("complete") == 7 and list(cov.values()).count("gap") == 1


def test_search_matches_pdf_p13(client):
    r = client.get("/api/search", params={"q": "collateral_value"}).json()["results"]
    assert [(x["kind"], x["name"], x["feeds_cells"], x["verb"]) for x in r] == [
        ("field", "collateral_master.collateral_value", 8, "Feeds"), ("field", "stg_collateral.market_value_eur", 8, "Feeds"),
        ("rule", "T7 · Credit risk mitigation, comprehensive method", 8, "Feeds"),
        ("mapping", "Collateral type → eligibility", 6, "Feeds"),
        ("control", "Collateral valuation age < 1 year", 8, "Guards")]
    assert r[4]["subtitle"] == "Hop 2 · waived on 30-Jun-2026 run"
    assert client.get("/api/search", params={"q": "collateral_value", "kind": "control"}).json()["results"][0]["uid"] == "control:H2_coll_age"


def test_impact_graph(client):
    r = client.get(f"/api/elements/{ELEMENT}/impact").json()
    assert r["stats"] == {"templates": 3, "cells": 8, "entities": 1, "dependency_levels": 2}
    assert len(r["impacted"]) == 8 and r["impacted"][0]["ref"] == "C 07.00 r0010 c0200"
    lanes = {n["label"]: n["lane"] for n in r["nodes"]}
    assert lanes["collateral_value"] == 0 and lanes["SA engine: CRM"] == 3 and lanes["C 03.00"] == 5
    assert {"src": ELEMENT, "dst": "field:stg_collateral.market_value_eur"} in r["edges"]
    assert client.get("/api/elements/field:nope.x/impact").status_code == 404


def test_what_if_reproduces_pdf_p15(client):
    r = client.post(f"/api/elements/{ELEMENT}/what-if", json={"haircut_pct": 15}).json()
    assert [x["change"] for x in r["rows"]] == [0.2, 0.2, 7.3, 1.5, 2.1, 1.6, 3.3]
    assert [x["from_element"] for x in r["rows"]] == [-1.2, -1.1, -48.5, -9.7, -14.2, -10.6, -21.4]
    assert [x["estimated"] for x in r["rows"]] == [1240.7, 1119.1, 619.6, 250.1, 907.9, 614.0, 9845.9]
    assert [x["pct_of_cell"] for x in r["rows"]] == [0.02, 0.02, 1.19, 0.6, 0.23, 0.26, 0.03]
    assert [x["threshold"] for x in r["rows"]] == ["within", "within", "above", "near", "within", "within", "within"]
    assert (r["ratio"]["reported_pct"], r["ratio"]["estimated_pct"], r["ratio"]["change_bps"]) == (14.82, 14.82, -0.5)
    assert r["decision"]["level"] == "discuss" and "1 cell(s) move more than 1%" in r["decision"]["text"]
    assert [(c["field"], c["cells"]) for c in r["cde_footprint"]][-1] == ("collateral_value", 8)
    big = client.post(f"/api/elements/{ELEMENT}/what-if", json={"haircut_pct": 100}).json()
    assert big["ratio"]["change_bps"] < -3
    assert client.post(f"/api/elements/{ELEMENT}/what-if", json={"haircut_pct": 150}).status_code == 400


def test_impact_report_is_pdf_and_audited(client, conn):
    r = client.post("/api/impact-report", json={"element_uid": ELEMENT, "haircut_pct": 15, "incident_ref": "INC-42"}, headers=SUP)
    assert r.status_code == 200 and r.content.startswith(b"%PDF") and b"INC-42" in r.content
    assert conn.execute("SELECT action FROM audit_log ORDER BY seq DESC LIMIT 1").fetchone()[0] == "exported impact report"


def test_regulator_view_and_snapshot(client, conn, st):
    before = conn.execute("SELECT count(*) FROM audit_log WHERE action = 'opened regulator view'").fetchone()[0]
    r = client.get(f"/api/cells/{CELL}/regulator", headers=SUP).json()
    assert [s["value"] for s in r["stages"]] == [1312.0, 1288.4, 1262.1, 1243.8, 1240.5]
    assert "(−23.6m)" in r["narrative"] and "(−26.3m)" in r["narrative"] and "(−18.3m" in r["narrative"] and "(−3.3m)" in r["narrative"]
    assert r["stats"] == {"hops": 7, "transformations": 8, "manual_adjustments": 1, "controls_passed": "13 / 14",
                          "unexplained": 0.0, "records": 3418}
    assert "1.2m" in r["limitations"][0]["text"] and "15-Dec-2026" in r["limitations"][0]["text"]
    assert r["evidence"]["last_pack"]["ref"] == "EP-2026-Q1-0017" and len(r["evidence"]["sections"]) == 10
    assert conn.execute("SELECT count(*) FROM audit_log WHERE action = 'opened regulator view'").fetchone()[0] == before + 1
    # snapshot is materialised once and immutable, even if the live facts change afterwards
    h1 = conn.execute("SELECT hash FROM lineage_snapshot WHERE cell_id = %s", (CELL,)).fetchone()[0]
    conn.execute("UPDATE hop_value SET amount = amount - 5 WHERE cell_id = %s AND step = 'fx'", (CELL,))
    try:
        assert snapshot.get(st, CELL, None)["waterfall"]["steps"][1]["amount"] == -11.7
        assert client.get(f"/api/cells/{CELL}/regulator", headers=SUP).json()["stages"][1]["value"] == 1288.4
        assert st.waterfall(CELL, None)["residual"] == 5.0  # live data did change
    finally:
        conn.execute("UPDATE hop_value SET amount = amount + 5 WHERE cell_id = %s AND step = 'fx'", (CELL,))
    assert conn.execute("SELECT hash FROM lineage_snapshot WHERE cell_id = %s", (CELL,)).fetchone()[0] == h1
    with pytest.raises(Exception, match="immutable"):
        conn.execute("UPDATE lineage_snapshot SET hash = 'x'")
    assert "•" in str(snapshot.get(st, CELL, None)["records"]["top"][0]) and "CPTY-00001107" not in str(snapshot.get(st, CELL, None))


def test_evidence_packs_pdf_and_docx(client, conn, tmp_path):
    from docx import Document

    a = client.post(f"/api/cells/{CELL}/evidence-pack", json={"sections": ALL, "format": "pdf"}, headers=SUP)
    b = client.post(f"/api/cells/{CELL}/evidence-pack", json={"sections": ["cover", "contributors"], "format": "docx"}, headers=SUP)
    assert a.status_code == b.status_code == 200
    ja, jb = a.json(), b.json()
    assert ja["ref"] != jb["ref"] and ja["ref"].startswith("EP-2026-Q2-")
    assert int(jb["ref"][-4:]) == int(ja["ref"][-4:]) + 1 and int(ja["ref"][-4:]) >= 42
    pdf = client.get(ja["download_url"])
    assert pdf.headers["content-type"] == "application/pdf" and pdf.content.startswith(b"%PDF")
    assert hashlib.sha256(pdf.content).hexdigest() == ja["sha256"] == conn.execute(
        "SELECT sha256 FROM evidence_pack WHERE ref = %s", (ja["ref"],)).fetchone()[0]
    assert ja["ref"].encode() in pdf.content and b"Supervisor" in pdf.content  # watermark: ref + viewer
    assert b"CPTY-00012764" not in pdf.content and b"2764" in pdf.content and b"LOS-FAC-00123458" not in pdf.content
    assert client.get(ja["download_url"]).content == pdf.content  # immutable: same bytes on every read
    doc = Document(io.BytesIO(client.get(jb["download_url"]).content))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "CPTY-•••2764" in text and "CPTY-00012764" not in text and "LOS-FAC-00••••58" in text
    assert jb["ref"] in doc.sections[0].header.paragraphs[0].text and "Supervisor [A]" in doc.sections[0].header.paragraphs[0].text
    assert (tmp_path / f"{ja['ref']}.pdf").exists()
    last = conn.execute("SELECT action, target FROM audit_log ORDER BY seq DESC LIMIT 1").fetchone()
    assert last == ("generated evidence pack", f"{jb['ref']} · C 07.00 r0010 c0200")
    assert client.post(f"/api/cells/{CELL}/evidence-pack", json={"sections": [], "format": "pdf"}, headers=SUP).status_code == 400
    assert client.post(f"/api/cells/{CELL}/evidence-pack", json={"sections": ALL, "format": "xls"}, headers=SUP).status_code == 400
    assert client.get("/api/evidence-packs/EP-2026-Q1-0017").status_code == 404  # seeded, archived
    assert client.get("/api/evidence-packs/EP-nope").status_code == 404
    reg = client.get(f"/api/cells/{CELL}/regulator", headers=SUP).json()["evidence"]
    assert reg["last_pack"]["ref"] == jb["ref"] and int(reg["next_ref"][-4:]) == int(jb["ref"][-4:]) + 1


def test_audit_chain_with_seeded_history(client, conn):
    r = client.get("/api/audit", params={"limit": 500}).json()
    assert r["chain_valid"] and r["retention"] == "7 years"
    actions = [e["action"] for e in reversed(r["entries"])][:6]
    assert actions == ["opened gap record", "changed metadata", "certified lineage", "certified lineage",
                       "exported drill-down (masked)", "revealed an identifier"]
    assert not any(e["target"].startswith("EP-2026-Q2-0042") for e in r["entries"] if e["ts"].startswith("2026-11"))
    assert audit.verify(conn)


def test_coverage_dashboard(client, conn):
    r = client.get("/api/coverage").json()
    k = r["kpis"]
    assert (k["cells_complete"], k["cells_note"], k["hops_certified_pct"], k["hops_note"], k["unexplained_breaks"],
            k["adjustments_approved"], k["open_gaps"], k["open_gaps_note"]) == (
        "7 / 10", "2 partial · 1 gap", 86, "43 of 50 hops", 0, "4 / 4", 1, "target 15-Dec-2026")
    assert k["evidence_packs"] == conn.execute("SELECT count(*) FROM evidence_pack WHERE ref LIKE 'EP-2026-Q2-%'").fetchone()[0] >= 6
    assert [t["status"] for t in r["by_template"]] == ["Complete", "Complete", "Complete", "Partial · 1 gap"]
    assert [(c["owner"], c["nodes"], c["pct"]) for c in r["certification"]][2] == ("Collateral data owner", 8, 63)


def test_adjustment_evidence_placeholder(client):
    r = client.get("/api/adjustments/ADJ-2026-Q2-014/evidence")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert client.get("/api/adjustments/nope/evidence").status_code == 404


def test_openlineage_event_is_idempotent(client, conn):
    ev = {"eventType": "COMPLETE", "eventTime": "2026-03-31T23:00:00Z", "run": {"runId": "run-ol-1", "facets": {
              "recordCounts": {"hop": 3, "in": 100, "out": 98, "submission_id": "2026-03-31.v1"},
              "controlResult": {"control_id": "OL_test_ctl", "name": "OL test control", "hop": 3, "status": "pass",
                                "detail": "ok", "submission_id": "2026-03-31.v1"},
              "manualAdjustment": {"id": "ADJ-OL-1", "amount": -1.5, "cell_id": "C07.00:Total:r0010:c0200", "maker": "[A]",
                                   "checker": "[B]", "submission_id": "2026-03-31.v1"}}},
          "job": {"namespace": "ol", "name": "load_ol_target"}, "inputs": [{"namespace": "ol", "name": "ol_src"}],
          "outputs": [{"namespace": "ol", "name": "ol_target", "facets": {"columnLineage": {"fields": {
              "total": {"inputFields": [{"namespace": "ol", "name": "ol_src", "field": "amount"}]}}}}}]}
    r1 = client.post("/api/ingest/openlineage", json=ev)
    assert r1.status_code == 200 and r1.json()["run_id"] == "run-ol-1"
    assert r1.json()["nodes_upserted"] == 4 and r1.json()["edges_upserted"] == 3
    assert db.cypher(conn, "MATCH (:Field {uid: 'field:ol_src.amount'})-[e:DERIVES]->(:Field {uid: 'field:ol_target.total'}) "
                           "RETURN e.job_id") == ["job:load_ol_target"]
    r2 = client.post("/api/ingest/openlineage", json=ev).json()
    assert r2["nodes_upserted"] == 0 and r2["edges_upserted"] == 0
    assert conn.execute("SELECT records_out FROM run WHERE run_id = 'run-ol-1'").fetchone()[0] == 98
    assert conn.execute("SELECT status FROM control_result WHERE control_id = 'OL_test_ctl'").fetchone()[0] == "pass"
    assert conn.execute("SELECT count(*) FROM adjustment WHERE id = 'ADJ-OL-1'").fetchone()[0] == 1
    assert client.post("/api/ingest/openlineage", json={"eventType": "START"}).status_code == 400


def test_ingest_sql_endpoint(client):
    r = client.post("/api/ingest/sql", json={"content": "INSERT INTO t_a (x) SELECT s.y AS x FROM t_b s;"}).json()
    assert r["edges_upserted"] == 3 and not r["errors"]  # 2 contains edges + 1 derives edge
    assert client.post("/api/ingest/sql", json={"content": "INSERT INTO t_a (x) SELECT s.y AS x FROM t_b s;"}).json()["edges_upserted"] == 0
    assert client.post("/api/ingest/sql", json={"path": "pyproject.toml"}).status_code == 400


def test_agent_role_cannot_read_raw_records(lineage_dsn):
    dsn = lineage_dsn.replace("user=lineage password=lineage", "user=lineage_agent password=lineage_agent")
    with psycopg.connect(dsn, autocommit=True) as c:
        with pytest.raises(psycopg.errors.InsufficientPrivilege):
            c.execute("SELECT counterparty FROM record_contribution LIMIT 1")
        with pytest.raises(psycopg.errors.InsufficientPrivilege):
            c.execute("SELECT * FROM audit_log LIMIT 1")
        rows = c.execute("SELECT counterparty, source_key FROM v_record_contribution_masked ORDER BY value DESC LIMIT 5").fetchall()
        assert rows[4] == ("CPTY-•••2764", "LOS-FAC-00••••58") and rows[1][1] == "CBS-ACC-7•••••02"
        with pytest.raises(psycopg.errors.ReadOnlySqlTransaction):
            c.execute("UPDATE cell_value SET value = 0")
